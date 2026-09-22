import csv
from datetime import date, datetime
from decimal import Decimal

import openpyxl

from app.exporters import export_invoice_to_csv, export_invoice_to_excel
from app.models import InvoiceData


def _valid_invoice(**overrides) -> InvoiceData:
    base = dict(
        invoice_number="FAC-2026-001",
        invoice_date=date(2026, 9, 7),
        total_amount=Decimal("1200"),
        currency="MAD",
        supplier_name="Atlas Trading Co.",
        customer_name="Jane Doe",
        customer_ICE="001234567000089",
        supplier_tax_id="001234567000000",
        subtotal=Decimal("1000"),
        tax_amount=Decimal("200"),
    )
    base.update(overrides)
    return InvoiceData(**base)
 
 
def test_csv_export_returns_path_and_creates_file(tmp_path):
    invoice = _valid_invoice()
    output_path = tmp_path / "invoice.csv"
 
    result = export_invoice_to_csv(invoice, output_path)
 
    assert result == output_path
    assert output_path.exists()
 
 
def test_export_csv_values_are_correct(tmp_path):
    invoice = _valid_invoice()
    output_path = tmp_path / "invoice.csv"
    export_invoice_to_csv(invoice, output_path)
 
    with output_path.open("r", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
 
    assert len(rows) == 1
    row = rows[0]
    assert row["invoice_date"] == "2026-09-07"
    assert row["subtotal"] == "1000"
    assert row["tax_amount"] == "200"
    assert row["total_amount"] == "1200"
    assert row["currency"] == "MAD"
    assert row["invoice_number"] == "FAC-2026-001"
    assert row["supplier_name"] == "Atlas Trading Co."
 
 
def test_export_none_fields_become_empty_cells(tmp_path):
    invoice = _valid_invoice(
        supplier_tax_id=None,
        customer_name=None,
        customer_ICE=None,
    )
    output_path = tmp_path / "invoice.csv"
    export_invoice_to_csv(invoice, output_path)
 
    with output_path.open("r", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
 
    row = rows[0]
    assert row["supplier_tax_id"] == ""
    assert row["customer_name"] == ""
    assert row["customer_ICE"] == ""
 
 
def test_csv_export_creates_missing_parent_directories(tmp_path):
    invoice = _valid_invoice()
    output_path = tmp_path / "exports" / "invoices" / "invoice.csv"
 
    assert not output_path.parent.exists()
 
    result = export_invoice_to_csv(invoice, output_path)
 
    assert result == output_path
    assert output_path.exists()


#===============================================================
# Excel Export Test
#=================================================================
def _valid_invoice(**overrides) -> InvoiceData:
    base = dict(
        invoice_number="FAC-2026-001",
        invoice_date=date(2026, 9, 7),
        total_amount=Decimal("1200"),
        currency="MAD",
        supplier_name="Atlas Trading Co.",
        customer_name="Jane Doe",
        customer_ICE="001234567000089",
        supplier_tax_id="001234567000000",
        subtotal=Decimal("1000"),
        tax_amount=Decimal("200"),
    )
    base.update(overrides)
    return InvoiceData(**base)
 
 
def test_xlsx_export_returns_path_and_creates_file(tmp_path):
    invoice = _valid_invoice()
    output_path = tmp_path / "invoice.xlsx"
 
    result = export_invoice_to_excel(invoice, output_path)
 
    assert result == output_path
    assert output_path.exists()
 
 
def test_workbook_contents_are_correct(tmp_path):
    invoice = _valid_invoice()
    output_path = tmp_path / "invoice.xlsx"
    export_invoice_to_excel(invoice, output_path)
 
    wb = openpyxl.load_workbook(output_path)
    assert "Invoice" in wb.sheetnames
    ws = wb["Invoice"]
 
    headers = [cell.value for cell in ws[1]]
    values = {
        headers[i]: ws.cell(row=2, column=i + 1).value
        for i in range(len(headers))
    }
 
    assert values["invoice_number"] == "FAC-2026-001"
    assert values["supplier_name"] == "Atlas Trading Co."
    assert values["subtotal"] == 1000
    assert values["tax_amount"] == 200
    assert values["total_amount"] == 1200
    assert values["currency"] == "MAD"
 
    # openpyxl may hand back a datetime rather than the original date,
    # depending on serialization -- normalize before comparing.
    invoice_date = values["invoice_date"]
    if isinstance(invoice_date, datetime):
        invoice_date = invoice_date.date()
    assert invoice_date == date(2026, 9, 7)
 
 
def test_optional_fields_remain_blank(tmp_path):
    invoice = _valid_invoice(
        supplier_tax_id=None,
        customer_name=None,
        customer_ICE=None,
    )
    output_path = tmp_path / "invoice.xlsx"
    export_invoice_to_excel(invoice, output_path)
 
    ws = openpyxl.load_workbook(output_path)["Invoice"]
    headers = [cell.value for cell in ws[1]]
    values = {
        headers[i]: ws.cell(row=2, column=i + 1).value
        for i in range(len(headers))
    }
 
    assert values["supplier_tax_id"] is None
    assert values["customer_name"] is None
    assert values["customer_ICE"] is None
 
 
def test_xlsx_export_creates_missing_parent_directories(tmp_path):
    invoice = _valid_invoice()
    output_path = tmp_path / "exports" / "invoices" / "invoice.xlsx"
 
    assert not output_path.parent.exists()
 
    result = export_invoice_to_excel(invoice, output_path)
 
    assert result == output_path
    assert output_path.exists()